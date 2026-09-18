"""pptx-toolkit 内部共享工具。

只依赖 Python 标准库与 python-pptx（MIT），不含任何授权受限的第三方代码。
所有脚本都会先调用 ensure_utf8_stdout()，避免在中文 Windows 的 GBK 控制台上
打印非 GBK 字符（例如 "→"、emoji、部分 CJK 扩展字符）时抛 UnicodeEncodeError。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

EXIT_OK = 0
EXIT_FAIL = 1
EXIT_USAGE = 2


def ensure_utf8_stdout() -> None:
    """把 stdout/stderr 切成 UTF-8，避免 GBK 控制台编码错误。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


def die(message: str, code: int = EXIT_FAIL) -> "None":
    print(f"[pptx-toolkit] 错误：{message}", file=sys.stderr)
    raise SystemExit(code)


def load_presentation(path: str):
    """打开 .pptx/.potx，失败时给出可读的中文错误。"""
    from pptx import Presentation

    p = Path(path)
    if not p.exists():
        die(f"文件不存在：{p}")
    if not p.is_file():
        die(f"不是文件：{p}")
    if p.suffix.lower() not in (".pptx", ".potx", ".pptm"):
        die(f"不是 PowerPoint 文件（.pptx/.potx/.pptm）：{p.name}")
    try:
        return Presentation(str(p))
    except Exception as exc:  # python-pptx 的异常类型比较杂，这里统一兜底
        die(
            f"无法打开 {p.name}：{type(exc).__name__}: {exc}\n"
            f"  提示：如果文件来自其它程序导出，先跑 check_deck.py 看包结构是否完整。"
        )


def iter_shapes(shapes):
    """深度遍历形状，组合形状（group）会被展开。"""
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) is not None:
            try:
                if shape.shape_type == 6 and hasattr(shape, "shapes"):  # MSO_SHAPE_TYPE.GROUP
                    yield from iter_shapes(shape.shapes)
            except (ValueError, AttributeError):
                continue


def shape_kind(shape) -> str:
    """给形状一个便于人读的类型名。"""
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return "text"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "shape_type", None) is not None:
        try:
            return str(shape.shape_type).split(" (")[0].lower().replace("mso_shape_type.", "")
        except (ValueError, AttributeError):
            return "unknown"
    return "unknown"


def shape_text(shape) -> str:
    """取形状的可读文本；没有文本帧时返回空串。"""
    if getattr(shape, "has_text_frame", False):
        return shape.text_frame.text
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            rows.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(rows)
    return ""


def notes_text(slide) -> str:
    """取演讲者备注；没有备注页时返回空串。"""
    try:
        if slide.has_notes_slide:
            return slide.notes_slide.notes_text_frame.text
    except (AttributeError, ValueError):
        pass
    return ""


def emu_to_inches(value) -> float:
    """EMU → 英寸，保留两位小数。"""
    try:
        return round(int(value) / 914400, 2)
    except (TypeError, ValueError):
        return 0.0


def emit_json(payload) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2))
